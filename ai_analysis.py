"""
Бесплатная автоматическая проверка загруженных файлов с прайсами.

Никаких платных API — вся логика работает локально:
1. Достаёт текст из файла (PDF, Excel, Word, CSV).
2. Ищет в тексте упоминания цен рядом с ключевыми словами по тем же
   категориям, что использует сайт (та же техника, что и в "Сканировать сайт").
3. Возвращает короткую сводку найденного — как подсказку для проверки,
   а не как гарантированно точный результат.

Фотографии прайс-листов автоматически не разбираются — для распознавания
текста на картинке нужен либо OCR (не всегда надёжен без дополнительных
системных зависимостей), либо платный ИИ. Файл в любом случае сохраняется
и доступен для просмотра вручную.
"""

import re

from database import SERVICE_TYPES
from scraper import KEYWORDS, PRICE_RE

IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "webp"}
MAX_SNIPPETS_PER_CATEGORY = 2


def extract_text(file_path, ext):
    """Возвращает текст из файла или None, если формат не поддерживается для извлечения текста."""
    ext = ext.lower()
    try:
        if ext in ("csv", "txt"):
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()

        if ext == "pdf":
            import pdfplumber
            parts = []
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages[:20]:  # ограничение на случай огромных файлов
                    text = page.extract_text() or ""
                    parts.append(text)
            return "\n".join(parts)

        if ext == "xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            parts = []
            for sheet in wb.worksheets[:5]:
                parts.append(f"--- Лист: {sheet.title} ---")
                for row in sheet.iter_rows(values_only=True, max_row=500):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        parts.append(" | ".join(cells))
            return "\n".join(parts)

        if ext == "docx":
            import docx
            doc = docx.Document(file_path)
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(parts)

    except Exception as e:
        return f"__EXTRACT_ERROR__: {e}"

    return None  # формат не поддержан (например .xls, .doc — старые бинарные форматы)


def analyze_price_file(company_name, file_path, ext):
    """
    Бесплатный локальный разбор файла. Возвращает (feedback_text, status), где
    status одно из: "done", "no_matches", "unsupported_image",
    "unsupported_format", "extract_error".
    """
    ext = ext.lower()

    if ext in IMAGE_EXTENSIONS:
        return (
            "Автоматическая проверка фотографий пока не поддерживается (это бесплатная версия "
            "без распознавания изображений). Файл сохранён — открой его и посмотри вручную.",
            "unsupported_image",
        )

    text = extract_text(file_path, ext)
    if text is None:
        return (
            f"Автоматическая проверка не поддерживается для формата .{ext} "
            "(старые форматы Excel/Word — .xls, .doc — читаются с ошибками). "
            "Файл сохранён, можно скачать и посмотреть вручную.",
            "unsupported_format",
        )
    if isinstance(text, str) and text.startswith("__EXTRACT_ERROR__"):
        return (f"Не удалось прочитать содержимое файла: {text.split(':', 1)[1].strip()}", "extract_error")
    if not text.strip():
        return ("Файл не содержит извлекаемого текста (возможно, это скан без текстового слоя).", "extract_error")

    results = {}
    for service_type, keywords in KEYWORDS.items():
        hits = []
        for kw in keywords:
            for match in re.finditer(kw, text, flags=re.IGNORECASE):
                start = max(0, match.start() - 50)
                end = min(len(text), match.end() + 50)
                snippet = text[start:end].replace("\n", " ").replace("|", " ").strip()
                snippet = re.sub(r"\s+", " ", snippet)
                if PRICE_RE.search(snippet):
                    hits.append(snippet)
        if hits:
            seen = []
            for h in hits:
                if h not in seen:
                    seen.append(h)
            results[service_type] = seen[:MAX_SNIPPETS_PER_CATEGORY]

    if not results:
        return (
            "Явных цен рядом с ключевыми словами не найдено. Возможно, тарифы оформлены "
            "таблицей без текстовых подписей — открой файл и посмотри вручную.",
            "no_matches",
        )

    lines = []
    for service_type, snippets in results.items():
        label = SERVICE_TYPES.get(service_type, service_type)
        joined = " | ".join(f"…{s}…" for s in snippets)
        lines.append(f"• {label}: {joined}")
    return ("\n".join(lines), "done")
