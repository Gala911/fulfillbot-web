"""
Backend-сервер (Flask) для веб-версии сравнения фулфилментов.

В отличие от чисто браузерной версии, здесь сервер сам делает сетевые запросы
(парсинг сайтов, отправка форм заявок) — поэтому CORS-ограничения браузера
здесь не действуют, ровно как в Telegram-боте.

Запуск:
    pip install -r requirements.txt
    python app.py
Открой http://localhost:5000 в браузере.
"""

from flask import Flask, jsonify, request, send_from_directory
import os
import uuid
from werkzeug.utils import secure_filename

import database as db
import forms
import scraper
import ai_analysis
from seed import seed, seed_extra

app = Flask(__name__, static_folder="static", static_url_path="")

UPLOAD_FOLDER = os.path.join(os.path.dirname(__file__), "uploads")
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
ALLOWED_EXTENSIONS = {"pdf", "xlsx", "xls", "csv", "doc", "docx", "png", "jpg", "jpeg", "webp"}
MAX_FILE_SIZE = 15 * 1024 * 1024  # 15 МБ
app.config["MAX_CONTENT_LENGTH"] = MAX_FILE_SIZE

db.init_db()
if not db.list_companies():
    seed()
seed_extra()  # безопасно на каждом запуске: добавляет только отсутствующие компании

DELETE_PASSCODE = os.environ.get("DELETE_PASSCODE")  # если не задан — удаление без пароля, только с подтверждением


@app.route("/")
def index():
    return send_from_directory(app.static_folder, "index.html")


# ---------------- companies ----------------

@app.route("/api/companies", methods=["GET"])
def api_list_companies():
    companies = db.list_companies()
    result = []
    for c in companies:
        tariffs = db.get_tariffs_for_company(c["id"])
        result.append({
            "id": c["id"],
            "name": c["name"],
            "website": c["website"],
            "city": c["city"],
            "notes": c["notes"],
            "price_file_url": f"/api/files/{c['price_file']}" if c["price_file"] else None,
            "price_file_name": c["price_file_original_name"],
            "ai_feedback": c["ai_feedback"],
            "ai_feedback_status": c["ai_feedback_status"],
            "partner_status": c["partner_status"],
            "user_rating": c["user_rating"],
            "user_comment": c["user_comment"],
            "tariffs": {
                t["service_type"]: {"price": t["price"], "unit": t["unit"], "comment": t["comment"]}
                for t in tariffs
            },
        })
    return jsonify(result)


def _allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def _save_uploaded_file(file_storage):
    """Сохраняет загруженный файл под уникальным именем, возвращает (stored_name, original_name).

    Важно: secure_filename() из werkzeug полностью вырезает нелатинские символы,
    поэтому файлы с русскими названиями (обычное дело для прайсов) после него
    превращаются в пустую строку. Расширение проверяем по исходному имени,
    для отображения берём только базовое имя без пути (без secure_filename),
    а на диске файл всё равно хранится под случайным UUID-именем — так что
    это безопасно.
    """
    raw_name = os.path.basename(file_storage.filename or "")
    if not raw_name or not _allowed_file(raw_name):
        return None, None
    ext = raw_name.rsplit(".", 1)[1].lower()
    stored_name = f"{uuid.uuid4().hex}.{ext}"
    file_storage.save(os.path.join(UPLOAD_FOLDER, stored_name))
    return stored_name, raw_name


def _run_ai_analysis(company_id, company_name, stored_name):
    """Синхронно анализирует только что загруженный файл и сохраняет результат в базу."""
    ext = stored_name.rsplit(".", 1)[1].lower()
    file_path = os.path.join(UPLOAD_FOLDER, stored_name)
    feedback, status = ai_analysis.analyze_price_file(company_name, file_path, ext)
    db.set_ai_feedback(company_id, feedback, status)
    return feedback, status


@app.route("/api/companies", methods=["POST"])
def api_add_company():
    # Поддерживаем и обычный JSON (без файла), и multipart/form-data (с файлом)
    if request.content_type and "multipart/form-data" in request.content_type:
        name = (request.form.get("name") or "").strip()
        website = (request.form.get("website") or "").strip() or None
        city = (request.form.get("city") or "").strip() or None
        uploaded = request.files.get("price_file")
    else:
        data = request.get_json(force=True)
        name = (data.get("name") or "").strip()
        website = (data.get("website") or "").strip() or None
        city = (data.get("city") or "").strip() or None
        uploaded = None

    if not name:
        return jsonify({"error": "Название обязательно"}), 400
    if db.get_company_by_name(name):
        return jsonify({"error": "Компания с таким названием уже есть"}), 409

    stored_name, original_name = (None, None)
    if uploaded and uploaded.filename:
        stored_name, original_name = _save_uploaded_file(uploaded)
        if stored_name is None:
            allowed = ", ".join(sorted(ALLOWED_EXTENSIONS))
            return jsonify({"error": f"Недопустимый формат файла. Разрешены: {allowed}"}), 400

    company_id = db.add_company(
        name=name, website=website, city=city,
        price_file=stored_name, price_file_original_name=original_name,
    )

    ai_feedback, ai_status = (None, None)
    if stored_name:
        ai_feedback, ai_status = _run_ai_analysis(company_id, name, stored_name)

    return jsonify({
        "id": company_id, "name": name, "website": website, "city": city,
        "price_file_url": f"/api/files/{stored_name}" if stored_name else None,
        "price_file_name": original_name,
        "ai_feedback": ai_feedback, "ai_feedback_status": ai_status,
    }), 201


@app.route("/api/companies/<int:company_id>/price_file", methods=["POST"])
def api_upload_price_file(company_id):
    """Прикрепить/заменить файл прайса у уже существующей компании."""
    company = db.get_company(company_id)
    if not company:
        return jsonify({"error": "Компания не найдена"}), 404
    uploaded = request.files.get("price_file")
    if not uploaded or not uploaded.filename:
        return jsonify({"error": "Файл не передан"}), 400
    stored_name, original_name = _save_uploaded_file(uploaded)
    if stored_name is None:
        allowed = ", ".join(sorted(ALLOWED_EXTENSIONS))
        return jsonify({"error": f"Недопустимый формат файла. Разрешены: {allowed}"}), 400
    # удаляем старый файл, если был
    if company["price_file"]:
        old_path = os.path.join(UPLOAD_FOLDER, company["price_file"])
        if os.path.exists(old_path):
            os.remove(old_path)
    db.set_price_file(company_id, stored_name, original_name)
    ai_feedback, ai_status = _run_ai_analysis(company_id, company["name"], stored_name)
    return jsonify({
        "price_file_url": f"/api/files/{stored_name}", "price_file_name": original_name,
        "ai_feedback": ai_feedback, "ai_feedback_status": ai_status,
    })


@app.route("/api/files/<path:filename>", methods=["GET"])
def api_get_file(filename):
    safe_name = secure_filename(filename)
    if safe_name != filename:
        return jsonify({"error": "Некорректное имя файла"}), 400
    return send_from_directory(UPLOAD_FOLDER, safe_name)


@app.route("/api/companies/<int:company_id>", methods=["DELETE"])
def api_delete_company(company_id):
    company = db.get_company(company_id)
    if not company:
        return jsonify({"error": "Не найдено"}), 404
    if DELETE_PASSCODE:
        provided = request.args.get("passcode") or (request.get_json(silent=True) or {}).get("passcode")
        if provided != DELETE_PASSCODE:
            return jsonify({"error": "Неверное кодовое слово"}), 403
    if company["price_file"]:
        file_path = os.path.join(UPLOAD_FOLDER, company["price_file"])
        if os.path.exists(file_path):
            os.remove(file_path)
    db.delete_company(company["name"])
    return jsonify({"ok": True})


@app.route("/api/companies/<int:company_id>/status", methods=["POST"])
def api_set_partner_status(company_id):
    company = db.get_company(company_id)
    if not company:
        return jsonify({"error": "Не найдено"}), 404
    data = request.get_json(force=True)
    status = data.get("status") or None
    try:
        db.set_partner_status(company_id, status)
    except ValueError:
        return jsonify({"error": "Неизвестный статус"}), 400
    return jsonify({"ok": True, "partner_status": status})


@app.route("/api/companies/<int:company_id>/review", methods=["POST"])
def api_set_review(company_id):
    company = db.get_company(company_id)
    if not company:
        return jsonify({"error": "Не найдено"}), 404
    data = request.get_json(force=True)
    rating = data.get("rating")
    comment = (data.get("comment") or "").strip() or None
    if rating is not None:
        try:
            rating = int(rating)
        except (TypeError, ValueError):
            return jsonify({"error": "Оценка должна быть числом от 1 до 5"}), 400
    try:
        db.set_review(company_id, rating=rating, comment=comment)
    except ValueError:
        return jsonify({"error": "Оценка должна быть от 1 до 5"}), 400
    return jsonify({"ok": True, "rating": rating, "comment": comment})


@app.route("/api/config", methods=["GET"])
def api_config():
    """Публичная информация о настройках сервера, нужная фронтенду (без секретов)."""
    return jsonify({
        "delete_requires_passcode": bool(DELETE_PASSCODE),
        "partner_statuses": db.PARTNER_STATUSES,
    })


@app.route("/api/companies/<int:company_id>", methods=["PATCH"])
def api_update_company(company_id):
    company = db.get_company(company_id)
    if not company:
        return jsonify({"error": "Не найдено"}), 404
    data = request.get_json(force=True)
    website = (data.get("website") or "").strip() or None
    city = (data.get("city") or "").strip() or None
    db.update_company_details(company_id, website=website, city=city)
    return jsonify({"ok": True, "website": website, "city": city})


# ---------------- tariffs ----------------

@app.route("/api/tariffs", methods=["POST"])
def api_set_tariff():
    data = request.get_json(force=True)
    company_id = data.get("company_id")
    service_type = data.get("service_type")
    price = data.get("price")
    unit = data.get("unit", "")
    comment = data.get("comment", "")

    if not company_id or service_type not in db.SERVICE_TYPES:
        return jsonify({"error": "Некорректные данные"}), 400
    try:
        price = float(price)
    except (TypeError, ValueError):
        return jsonify({"error": "Цена должна быть числом"}), 400

    db.set_tariff(company_id, service_type, price=price, unit=unit, comment=comment)
    return jsonify({"ok": True})


@app.route("/api/companies/<int:company_id>/history", methods=["GET"])
def api_tariff_history(company_id):
    rows = db.get_tariff_history(company_id)
    return jsonify([
        {
            "service_type": r["service_type"],
            "service_label": db.SERVICE_TYPES.get(r["service_type"], r["service_type"]),
            "old_price": r["old_price"], "old_unit": r["old_unit"], "old_comment": r["old_comment"],
            "changed_at": r["changed_at"],
        }
        for r in rows
    ])


@app.route("/api/services", methods=["GET"])
def api_services():
    return jsonify(db.SERVICE_TYPES)


@app.route("/api/compare/<service_type>", methods=["GET"])
def api_compare(service_type):
    if service_type not in db.SERVICE_TYPES:
        return jsonify({"error": "Неизвестный вид услуги"}), 400
    rows = db.compare_service(service_type)
    return jsonify([
        {"name": r["name"], "website": r["website"], "city": r["city"], "price": r["price"],
         "unit": r["unit"], "comment": r["comment"], "updated_at": r["updated_at"]}
        for r in rows
    ])


# ---------------- scan (best-effort price discovery) ----------------

@app.route("/api/scan/<int:company_id>", methods=["POST"])
def api_scan(company_id):
    company = db.get_company(company_id)
    if not company:
        return jsonify({"error": "Не найдено"}), 404
    if not company["website"]:
        return jsonify({"error": "У компании не указан сайт"}), 400
    results = scraper.find_candidate_prices(company["website"])
    if "_error" in results:
        return jsonify({"error": results["_error"]}), 502
    readable = {db.SERVICE_TYPES[k]: v for k, v in results.items()}
    return jsonify({"candidates": readable})


# ---------------- send request (best-effort form autofill) ----------------

@app.route("/api/send_request", methods=["POST"])
def api_send_request():
    data = request.get_json(force=True)
    company_id = data.get("company_id")
    lead = {
        "name": data.get("name", ""),
        "phone": data.get("phone", ""),
        "email": data.get("email", ""),
        "message": data.get("message", ""),
    }
    confirm = bool(data.get("confirm", False))

    company = db.get_company(company_id)
    if not company:
        return jsonify({"error": "Компания не найдена"}), 404
    if not lead["name"] or not lead["phone"]:
        return jsonify({"error": "Укажите имя и телефон"}), 400

    if not company["website"]:
        text = forms.build_manual_request_text(lead, company["name"])
        db.log_request(company_id, "manual_no_website", text)
        return jsonify({"mode": "manual", "text": text})

    ok, details = forms.find_and_submit_form(company["website"], lead, dry_run=not confirm)
    if not ok:
        text = forms.build_manual_request_text(lead, company["name"])
        db.log_request(company_id, "manual_form_failed", details)
        return jsonify({"mode": "manual", "text": text, "reason": details})

    if confirm:
        db.log_request(company_id, "sent", details)
        return jsonify({"mode": "sent", "details": details})
    else:
        return jsonify({"mode": "preview", "details": details})


@app.route("/api/export/excel", methods=["GET"])
def api_export_excel():
    """Экспорт полного сравнения (все компании × все услуги) в один Excel-файл."""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    companies = db.list_companies()
    wb = Workbook()
    ws = wb.active
    ws.title = "Сравнение тарифов"

    headers = ["Компания", "Город", "Сайт"] + list(db.SERVICE_TYPES.values())
    ws.append(headers)
    header_fill = PatternFill(start_color="EDE7F9", end_color="EDE7F9", fill_type="solid")
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for c in companies:
        tariffs = {t["service_type"]: t for t in db.get_tariffs_for_company(c["id"])}
        row = [c["name"], c["city"] or "", c["website"] or ""]
        for service_key in db.SERVICE_TYPES:
            t = tariffs.get(service_key)
            if t and t["price"] is not None:
                cell_value = f"{t['price']} {t['unit'] or ''}".strip()
                if t["comment"]:
                    cell_value += f" ({t['comment']})"
            else:
                cell_value = ""
            row.append(cell_value)
        ws.append(row)

    # автоширина колонок (в разумных пределах)
    for col_idx, header in enumerate(headers, start=1):
        letter = get_column_letter(col_idx)
        max_len = max([len(header)] + [len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(2, ws.max_row + 1)])
        ws.column_dimensions[letter].width = min(max(12, max_len + 2), 45)
    ws.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    from flask import send_file
    return send_file(
        buffer,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="сравнение_фулфилментов.xlsx",
    )


@app.errorhandler(413)
def too_large(e):
    mb = MAX_FILE_SIZE // (1024 * 1024)
    return jsonify({"error": f"Файл слишком большой (максимум {mb} МБ)"}), 413


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=False)
